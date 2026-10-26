"""Unit tests for Excel workbook export (Milestone 6).

Verifies multi-sheet workbook generation, openpyxl compatibility, and data integrity.
"""

import io
import unittest
import openpyxl

from core.exporter import export_ledger_to_excel


class TestExcelExport(unittest.TestCase):
    """Verifies generation of valid multi-sheet .xlsx files from confirmed records."""

    def setUp(self) -> None:
        self.sample_entries = [
            {
                "id": "e-1",
                "customer_name": "Ramesh Kumar",
                "customer_phone": "9876543210",
                "item_name": "Toor Dal",
                "standardized_item_name": "Toor Dal",
                "quantity": 2.0,
                "unit": "kg",
                "unit_price": 140.0,
                "total_amount": 280.0,
                "paid_amount": 100.0,
                "due_amount": 180.0,
                "transaction_type": "credit",
                "status": "confirmed",
                "entry_date": "2026-10-20",
                "batch_id": "batch-101",
            },
            {
                "id": "e-2",
                "customer_name": "Suresh Gowda",
                "customer_phone": "9900112233",
                "item_name": "Cooking Oil",
                "standardized_item_name": "Cooking Oil",
                "quantity": 1.0,
                "unit": "litre",
                "unit_price": 150.0,
                "total_amount": 150.0,
                "paid_amount": 150.0,
                "due_amount": 0.0,
                "transaction_type": "cash",
                "status": "confirmed",
                "entry_date": "2026-10-22",
                "batch_id": "batch-101",
            },
        ]

    def test_export_ledger_to_excel_valid_data(self) -> None:
        """Generates valid .xlsx binary containing all three required sheets."""
        excel_bytes = export_ledger_to_excel(self.sample_entries, shop_name="Test Store")
        self.assertTrue(len(excel_bytes) > 0)
        # Check ZIP / XLSX magic header
        self.assertTrue(excel_bytes.startswith(b"PK\x03\x04"))

        # Verify openpyxl can load the workbook
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        expected_sheets = ["Transactions", "Customer Dues", "Product Performance"]
        self.assertEqual(wb.sheetnames, expected_sheets)

        # Check Transactions sheet has 2 rows + 1 header = 3 rows
        ws_tx = wb["Transactions"]
        self.assertEqual(ws_tx.max_row, 3)
        self.assertEqual(ws_tx.cell(row=1, column=1).value, "Date")
        self.assertEqual(ws_tx.cell(row=2, column=2).value, "Ramesh Kumar")

        # Check Customer Dues sheet
        ws_cust = wb["Customer Dues"]
        self.assertEqual(ws_cust.max_row, 3)  # 2 customers + 1 header
        self.assertEqual(ws_cust.cell(row=1, column=1).value, "Customer Name")

        # Check Product Performance sheet
        ws_prod = wb["Product Performance"]
        self.assertEqual(ws_prod.max_row, 3)  # 2 products + 1 header
        self.assertEqual(ws_prod.cell(row=1, column=1).value, "Product Name")

    def test_export_ledger_to_excel_empty_data(self) -> None:
        """Handles empty entries list cleanly without raising exceptions."""
        excel_bytes = export_ledger_to_excel([], shop_name="Empty Store")
        self.assertTrue(len(excel_bytes) > 0)
        self.assertTrue(excel_bytes.startswith(b"PK\x03\x04"))

        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        self.assertEqual(wb.sheetnames, ["Transactions", "Customer Dues", "Product Performance"])

    def test_export_ledger_to_excel_null_values(self) -> None:
        """Handles entries with None for customer name, paid amount, and due amount."""
        entries_with_nulls = [
            {
                "customer_name": None,
                "item_name": "Matchbox",
                "quantity": 10.0,
                "total_amount": 20.0,
                "paid_amount": None,
                "due_amount": None,
            }
        ]
        excel_bytes = export_ledger_to_excel(entries_with_nulls, shop_name="Corner Kirana")
        self.assertTrue(len(excel_bytes) > 0)
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        ws_tx = wb["Transactions"]
        self.assertEqual(ws_tx.cell(row=2, column=4).value, "Matchbox")
        self.assertIn(ws_tx.cell(row=2, column=2).value, [None, ""])  # Customer name is empty / None

    def test_export_customer_dues_calculation_matches(self) -> None:
        """Verifies customer dues in exported sheet match expected balances."""
        excel_bytes = export_ledger_to_excel(self.sample_entries, shop_name="Test Store")
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        ws_cust = wb["Customer Dues"]
        # Ramesh Kumar had total_billed=280, paid=100, due=180
        # Suresh Gowda had total_billed=150, paid=150, due=0
        # Row 2 should be Ramesh (higher due first)
        self.assertEqual(ws_cust.cell(row=2, column=1).value, "Ramesh Kumar")
        self.assertEqual(float(ws_cust.cell(row=2, column=5).value), 180.0)
        self.assertEqual(ws_cust.cell(row=3, column=1).value, "Suresh Gowda")
        self.assertEqual(float(ws_cust.cell(row=3, column=5).value), 0.0)


if __name__ == "__main__":
    unittest.main()
